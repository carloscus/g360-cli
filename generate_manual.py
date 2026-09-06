# -*- coding: utf-8 -*-
"""Manual PPTX de G360 CLI — inspirado en ventas-pulse, mais colorido e vibrante.
Widescreen 16:9 (13.333 x 7.5 in), branding G360/CIPSA com paleta completa."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ============================================
# PALETA G360 MULTI-COLOR (Asiento/Corporate)
# ============================================
C = {
    # Verde principal (asiento)
    'primary': RGBColor(0x00, 0xD0, 0x84),
    'primary_dark': RGBColor(0x00, 0xA8, 0x6B),
    'primary_light': RGBColor(0x33, 0xE6, 0xAA),
    
    # Azul secundario
    'accent': RGBColor(0x0E, 0x74, 0x90),
    'accent_light': RGBColor(0x22, 0xB8, 0xDE),
    
    # Amarelo/Warning
    'warning': RGBColor(0xF5, 0x9E, 0x0B),
    'warning_light': RGBColor(0xFB, 0xBD, 0x24),
    
    # Vermelho/Danger
    'danger': RGBColor(0xEF, 0x44, 0x44),
    'danger_light': RGBColor(0xF8, 0x71, 0x71),
    
    # Roxo/Debug
    'purple': RGBColor(0x8B, 0x5C, 0xF6),
    'purple_light': RGBColor(0xA7, 0x8B, 0xF1),
    
    # Fundos
    'white': RGBColor(0xFF, 0xFF, 0xFF),
    'bg': RGBColor(0xF8, 0xFA, 0xFC),
    'surface': RGBColor(0xF1, 0xF5, 0xF9),
    'surface_dark': RGBColor(0xE2, 0xE8, 0xF0),
    
    # Textos
    'text': RGBColor(0x0F, 0x17, 0x2A),
    'text_secondary': RGBColor(0x47, 0x55, 0x69),
    'muted': RGBColor(0x94, 0xA3, 0xB8),
    
    # Bordas
    'border': RGBColor(0xCB, 0xD4, 0xDE),
    'border_light': RGBColor(0xE2, 0xE8, 0xF0),
}

W, H = 13.333, 7.5
TOTAL = 10
CW = W - 1.2  # conteudo width

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

def nueva():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = C['bg']
    bg.line.fill.background()
    return s

def tb(s, x, y, w, h):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return box.text_frame

def p(tf, text, size=14, bold=False, color=None, first=False, space=6, align=PP_ALIGN.LEFT):
    par = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    par.space_after = Pt(space)
    par.alignment = align
    r = par.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold
    f.color.rgb = color or C['text']; f.name = 'Segoe UI'
    return par

def secao(s, titulo, num):
    # Header bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(1.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = C['primary']
    bar.line.fill.background()
    
    # Number badge
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.25), Inches(0.6), Inches(0.6))
    badge.fill.solid(); badge.fill.fore_color.rgb = C['white']
    badge.line.fill.background()
    btf = badge.text_frame; btf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = btf.paragraphs[0].add_run(); r.text = f'{num:02d}'
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = C['primary']
    
    # Title
    ttf = tb(s, 1.2, 0.28, 10, 0.6)
    p(ttf, titulo, 24, True, C['white'], first=True)
    
    # Footer
    ft = tb(s, 10.5, 7.0, 2.5, 0.35)
    ft.paragraphs[0].alignment = PP_ALIGN.RIGHT
    r = ft.paragraphs[0].add_run(); r.text = f'G360 CLI · v1.17 · {num}/{TOTAL}'
    r.font.size = Pt(9); r.font.color.rgb = C['muted']; r.font.name = 'Segoe UI'

def card(s, x, y, w, h, titulo, corpo, icon='•', icolor=C['primary']):
    # Card background
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = C['white']
    card.line.color.rgb = C['border']; card.line.width = Pt(1)
    # Top accent
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.08))
    accent.fill.solid(); accent.fill.fore_color.rgb = icolor
    accent.line.fill.background()
    # Icon circle
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(y + 0.25), Inches(0.35), Inches(0.35))
    circ.fill.solid(); circ.fill.fore_color.rgb = icolor; circ.line.fill.background()
    itf = circ.text_frame; itf.paragraphs[0].alignment = PP_ALIGN.CENTER
    ir = itf.paragraphs[0].add_run(); ir.text = icon
    ir.font.size = Pt(14); ir.font.color.rgb = C['white']
    # Text
    tf = tb(s, x + 0.65, y + 0.2, w - 0.8, h - 0.3)
    tf.word_wrap = True
    p(tf, titulo, 13, True, C['text'], first=True, space=4)
    for line in corpo:
        p(tf, line, 10.5, False, C['text_secondary'], space=2)

def row_card(s, x, y, w, h, cols):
    """Row card with multiple colored columns"""
    cw = w / len(cols)
    for i, (color, titulo, corpo) in enumerate(cols):
        cx = x + i * cw
        # Column bg
        col = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(y), Inches(cw - 0.05), Inches(h))
        col.fill.solid(); col.fill.fore_color.rgb = color
        col.line.fill.background()
        # Text
        tf = tb(s, cx + 0.1, y + 0.15, cw - 0.2, h - 0.25)
        tf.word_wrap = True
        p(tf, titulo, 11, True, C['white'], first=True, space=3)
        for line in corpo:
            p(tf, line, 9, False, C['white'], space=1)

def placeholder(s, x, y, w, h, arquivo, etiqueta):
    from pathlib import Path
    shots = Path('assets/screenshots')
    if (shots / arquivo).exists():
        try:
            from PIL import Image
            img = Image.open(shots / arquivo)
            iw, ih = img.size
            ir = iw / ih
            fr = w / h
            if ir > fr: dw, dh = w, w / ir
            else: dh, dw = h, h * ir
            dx = x + (w - dw) / 2
            dy = y + (h - dh) / 2
            s.shapes.add_picture(str((shots / arquivo).resolve()), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
            return
        except: pass
    ph = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    ph.fill.solid(); ph.fill.fore_color.rgb = C['surface']
    ph.line.color.rgb = C['primary']; ph.line.width = Pt(1.5); ph.line.dash_style = 4
    tf = ph.text_frame; tf.word_wrap = True
    p(tf, f'[ {etiqueta.upper()} ]', 11, True, C['muted'], first=True, align=PP_ALIGN.CENTER)
    p(tf, f'Salvar: assets/screenshots/{arquivo}', 9, False, C['muted'], align=PP_ALIGN.CENTER)

# ===== 1. PORTADA =====
s = nueva()
# Gradient-like top bar
top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(2.8))
top.fill.solid(); top.fill.fore_color.rgb = C['primary']
top.line.fill.background()
# Accent stripe
stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.8), Inches(W), Inches(0.15))
stripe.fill.solid(); stripe.fill.fore_color.rgb = C['accent']
stripe.line.fill.background()
# Title area
ttf = tb(s, 1, 0.8, 11.3, 1.2)
p(ttf, 'G360 CLI', 52, True, C['white'], first=True, align=PP_ALIGN.CENTER)
stf = tb(s, 1.5, 2.0, 10.3, 0.6)
p(stf, 'Scaffolding & Automation para Ecosistema G360', 18, False, C['white'], align=PP_ALIGN.CENTER)
# Content cards
card(s, 1.5, 3.2, 3.5, 1.4, 'Boostrapping', ['Projet padrao G360','Templates web/python','Brand system integrado'], '🚀', C['primary'])
card(s, 5.3, 3.2, 3.5, 1.4, 'Ingestao ERP', ['Normalizacao .xls/.xlsx','Commercial engine','Classificacao docs'], '📊', C['accent'])
card(s, 9.1, 3.2, 3.5, 1.4, 'Auditoria', ['Compliance G360','Validacao automatica','Lint de nomenclatura'], '🔍', C['warning'])
# Footer info
ift = tb(s, 1, 5.0, 11.3, 1.5)
p(ift, 'Ferramenta de linha de comando para iniciacao rapida de projetos', 14, False, C['text_secondary'], first=True, align=PP_ALIGN.CENTER)
p(ift, 'Node.js >= 18 · Distribuicao npm global · 15+ comandos', 12, False, C['muted'], align=PP_ALIGN.CENTER, space=4)
p(ift, 'v1.17.0 · powered by G360', 11, True, C['primary'], align=PP_ALIGN.CENTER)

# ===== 2. INTRO =====
s = nueva(); secao(s, 'INTRODUCAO', 2)
# What is it
itf = tb(s, 0.6, 1.3, 12.1, 0.8)
p(itf, 'O que e o G360 CLI?', 18, True, C['text'], first=True, space=4)
ctf = tb(s, 0.6, 2.1, 12.1, 1.2)
p(ctf, 'CLI tool para bootstrap de projetos G360 com estrutura padrao, assets embebidos, e compliance automatico. Facilita a criacao de apps web (Lit, Svelte, Solid), desktop (Flet, CustomTkinter), e scripts Python com identidade G360 integrada.', 12, False, C['text_secondary'])
# Key features row
row_card(s, 0.6, 3.5, 12.1, 1.6, [
    (C['primary'], 'Templates', ['10+ templates\nweb & desktop', 'Auto-detecta skill']),
    (C['accent'], 'Assets', ['Brand system v2\nLogos, favicons,\nicons PWA']),
    (C['warning'], 'Ingestao', ['Normaliza ERP\n.xls/.xlsx/.csv', 'Commercial engine']),
    (C['purple'], 'Auditoria', ['Verifica compliance\nValida estrutura', 'Lint automatico']),
])
# Bottom note
ntf = tb(s, 0.6, 5.4, 12.1, 1.5)
p(ntf, 'Instalacao', 14, True, C['text'], first=True, space=3)
p(ntf, 'npm install -g g360-cli', 12, False, C['primary'], space=2)
p(ntf, 'g360 init meu-projeto --template lit-web --skill corporativo-movil', 11, False, C['text_secondary'])

# ===== 3. COMANDOS =====
s = nueva(); secao(s, 'COMANDOS PRINCIPAIS', 3)
cmds = [
    ('init', 'Criar projeto', C['primary']),
    ('bring', 'Trazer assets', C['accent']),
    ('audit', 'Auditar projeto', C['warning']),
    ('clean', 'Limpar codigo', C['danger']),
    ('pptx', 'Gerar manual', C['purple']),
    ('convert', 'Converter projeto', C['primary']),
    ('signature', 'Instalar marca', C['accent']),
    ('ingest', 'Processar ERP', C['warning']),
]
y = 1.2
for i, (cmd, desc, color) in enumerate(cmds):
    x = 0.6 + (i % 4) * 3.1
    if i >= 4: y += 1.8
    # Command box
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.9), Inches(1.5))
    box.fill.solid(); box.fill.fore_color.rgb = C['white']
    box.line.color.rgb = color; box.line.width = Pt(2)
    # Accent top
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2.9), Inches(0.08))
    acc.fill.solid(); acc.fill.fore_color.rgb = color
    acc.line.fill.background()
    # Code
    ctf = tb(s, x + 0.15, y + 0.25, 2.6, 0.5)
    p(ctf, f'g360 {cmd}', 13, True, color, first=True)
    # Desc
    dtf = tb(s, x + 0.15, y + 0.75, 2.6, 0.6)
    p(dtf, desc, 10, False, C['text_secondary'], first=True)

# ===== 4. TEMPLATES =====
s = nueva(); secao(s, 'TEMPLATES DISPONIVEIS', 4)
templates = [
    ('Web PWA', 'React, Lit, Svelte, Solid', 'PWA instalavel com Service Worker'),
    ('Python Flet', 'Flet desktop (Padrao)', 'App desktop com tema G360'),
    ('Python CLI', 'Script CLI argparse', 'Ferramenta de linha de comando'),
    ('CustomTkinter', 'Desktop moderno', 'GUI escura com CTk'),
]
y = 1.3
for i, (name, stack, desc) in enumerate(templates):
    x = 0.6 + (i % 2) * 6.2
    if i >= 2: y += 2.0
    card(s, x, y, 5.9, 1.7, name, [f'Technology: {stack}', f'Desc: {desc}'], '⬡', C['primary'] if i%2==0 else C['accent'])
# Full width note
ntf = tb(s, 0.6, 5.5, 12.1, 1.2)
p(ntf, 'Todos os templates incluem: G360 signature, brand assets, skill.json, e estrututa padrao de pastas (src/core, src/ui, assets/).', 11, False, C['text_secondary'])

# ===== 5. BRAND SYSTEM =====
s = nueva(); secao(s, 'BRAND SYSTEM v2.0', 5)
# Color palette display
colors = [
    (C['primary'], 'Verde\nAsiento'),
    (C['accent'], 'Azul\nSecundario'),
    (C['warning'], 'Amarelo\nWarning'),
    (C['danger'], 'Vermelho\nDanger'),
    (C['purple'], 'Roxo\nDebug'),
]
x = 0.6
for color, label in colors:
    swatch = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.4), Inches(2.0), Inches(1.4))
    swatch.fill.solid(); swatch.fill.fore_color.rgb = color
    swatch.line.fill.background()
    ltf = tb(s, x, 2.8, 2.0, 0.6)
    p(ltf, label, 10, True, color, first=True, align=PP_ALIGN.CENTER)
    x += 2.2
# Brand elements
brand_items = [
    ('Isotipo', '3 pontos verticais + chevron >'),
    ('Logotipos', 'G360 light/dark, CIPSA solid/borde'),
    ('Favicons', '16x16, 32x32, apple-touch, ICO'),
    ('PWA Icons', '192x192, 512x512, transparent'),
]
y = 3.3
for titulo, desc in brand_items:
    tf = tb(s, 0.6, y, 12.1, 0.5)
    p(tf, f'▸ {titulo}: {desc}', 11, False, C['text_secondary'], first=True)
    y += 0.55

# ===== 6. WORKFLOW =====
s = nueva(); secao(s, 'WORKFLOW DE USO', 6)
steps = [
    ('1', 'init', 'g360 init projeto --template lit-web --skill corporativo-movil'),
    ('2', 'bring', 'g360 bring (assets, brand, components)'),
    ('3', 'develop', 'Desenvolver com padroes G360'),
    ('4', 'audit', 'g360 audit (verificar compliance)'),
    ('5', 'deploy', 'g360 pptx + deploy GitHub Pages'),
]
y = 1.3
for num, cmd, desc in steps:
    # Number circle
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(y), Inches(0.5), Inches(0.5))
    circ.fill.solid(); circ.fill.fore_color.rgb = C['primary']
    circ.line.fill.background()
    ntf = circ.text_frame; ntf.paragraphs[0].alignment = PP_ALIGN.CENTER
    nr = ntf.paragraphs[0].add_run(); nr.text = num
    nr.font.size = Pt(16); nr.font.bold = True; nr.font.color.rgb = C['white']
    # Command
    ctf = tb(s, 1.3, y - 0.05, 6, 0.6)
    p(ctf, f'{cmd}: {desc}', 12, False, C['text'], first=True)
    y += 1.0
# Arrow connectors
for i in range(4):
    arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(0.8), Inches(1.85 + i * 1.0), Inches(0.2), Inches(0.25))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = C['primary']
    arrow.line.fill.background()
# Screenshot placeholder
placeholder(s, 7.8, 1.2, 5.0, 4.8, 'cli-terminal.png', 'Terminal com g360 init em acao')

# ===== 7. INGESTA ERP =====
s = nueva(); secao(s, 'INGESTA ERP & COMMERCIAL ENGINE', 7)
# Left: Ingestion
card(s, 0.6, 1.3, 5.8, 2.4, 'Normalizacao ERP', [
    'Suporta: .xls, .xlsx, .csv',
    'Auto-detecta: SAP, StarSoft, Spring',
    'Parseia referencias: F01/201-243065',
    'Separa sucursais, normaliza monetario',
    'Classifica: VENDA, DEVOLUCAO, AJUSTE',
], '📄', C['accent'])
# Right: Engine
card(s, 6.8, 1.3, 5.8, 2.4, 'Commercial Engine', [
    'classify_base(): VENDA/DEVOLUCAO/AJUSTE',
    'build_invoice_index(): index para cruce',
    'resolve_document_relationships()',
    'calculate_prices(): PRECO_BASE, RECARGO',
    'SUBTIPO_AJUSTE: PRECO_LINEA, CARGO_FIJO',
], '⚙️', C['purple'])
# Bottom flow
flow = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.0), Inches(12.1), Inches(2.8))
flow.fill.solid(); flow.fill.fore_color.rgb = C['surface']
flow.line.color.rgb = C['border']; flow.line.width = Pt(1)
# Steps
flow_steps = [
    ('Input', 'Arquivo ERP (.xls/.xlsx)', C['primary']),
    ('Parse', 'Leitura e normalizacao', C['accent']),
    ('Classify', 'Commercial engine', C['purple']),
    ('Output', 'CSV maestro normalizado', C['warning']),
]
x = 0.9
for nome, desc, cor in flow_steps:
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.3), Inches(2.7), Inches(2.1))
    box.fill.solid(); box.fill.fore_color.rgb = C['white']
    box.line.color.rgb = cor; box.line.width = Pt(2)
    # Title
    btf = tb(s, x + 0.1, 4.4, 2.5, 0.5)
    p(btf, nome, 13, True, cor, first=True, align=PP_ALIGN.CENTER)
    # Desc
    dtf = tb(s, x + 0.1, 5.0, 2.5, 1.2)
    p(dtf, desc, 10, False, C['text_secondary'], first=True, align=PP_ALIGN.CENTER)
    x += 2.9

# ===== 8. SKILLS =====
s = nueva(); secao(s, 'SKILLS & CONFIGURACAO', 8)
skills = [
    ('corporativo', 'PC clientes', C['primary']),
    ('corporativo-movil', 'Mobile clientes', C['accent']),
    ('moderno', 'PC proprias', C['warning']),
    ('moderno-movil', 'Mobile proprias', C['purple']),
    ('minimalista', 'Scripts/CLI', C['danger']),
    ('cipsa', 'Marca CIPSA', RGBColor(0xC4, 0x1E, 0x3A)),
    ('flet-desktop', 'Desktop Flet', C['accent']),
    ('lit-web', 'Lit Components', C['primary']),
]
x, y = 0.6, 1.3
for nome, desc, cor in skills:
    card(s, x, y, 2.8, 1.2, nome, [desc], '⬡', cor)
    x += 3.0
    if x > 10: x = 0.6; y += 1.5
# Config note
ctf = tb(s, 0.6, 5.5, 12.1, 1.2)
p(ctf, 'Configuracao global em ~/.g360/ com cache de assets em ~/.g360/cache/. Use g360 config --list para ver opcoes.', 11, False, C['text_secondary'])

# ===== 9. LIMITACOES =====
s = nueva(); secao(s, 'LIMITACOES CONHECIDAS', 9)
lims = [
    ('Python requerido', 'Alguns comandos exigem Python 3.11+ (commercial_engine, ingestion)'),
    ('Assets embebidos', 'Templates grandes aumentam tamanho do pacote npm (~50MB)'),
    ('Conflito de nomes', 'Se projeto ja existir sem --force, Comando falha'),
    ('Validate rigido', 'Audit pode ser muito strict para projetos legacy'),
]
y = 1.3
for titulo, desc in lims:
    # Warning icon
    warn = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.6), Inches(y), Inches(0.4), Inches(0.4))
    warn.fill.solid(); warn.fill.fore_color.rgb = C['warning']
    warn.line.fill.background()
    wtf = warn.text_frame; wtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    wr = wtf.paragraphs[0].add_run(); wr.text = '!'
    wr.font.size = Pt(14); wr.font.bold = True; wr.font.color.rgb = C['white']
    # Text
    ttf = tb(s, 1.2, y - 0.05, 11.5, 0.85)
    p(ttf, f'{titulo}: {desc}', 11, False, C['text_secondary'], first=True)
    y += 0.95

# ===== 10. RESUMO =====
s = nueva(); secao(s, 'RESUMO & PROXIMOS PASSOS', 10)
# Summary box
summary = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.3), Inches(12.1), Inches(1.5))
summary.fill.solid(); summary.fill.fore_color.rgb = C['primary']
summary.line.fill.background()
stf = tb(s, 0.8, 1.4, 11.7, 1.3)
p(stf, 'G360 CLI: Padronizacao e velocidade para o ecosistema', 16, True, C['white'], first=True, align=PP_ALIGN.CENTER, space=4)
p(stf, 'Iniciacao rapida · Brand system unificado · Ingestao ERP · Auditoria automatica · Geracao de manuais', 12, False, C['white'], align=PP_ALIGN.CENTER)
# Checklist
checks = [
    '10+ templates para web (Lit, Svelte, Solid, React) e desktop (Flet, CTk)',
    'Brand system v2 com logos, favicons, e icons PWA prontos',
    'Commercial engine para classificacao de documentos ERP',
    'Auditoria automatica de compliance G360',
    'Geracao de manuais PPTX a partir doCodigo',
    'Disponivel como pacote npm global (g360) e PyPI (g360-core)',
]
y = 3.1
for check in checks:
    tf = tb(s, 1.0, y, 11.5, 0.45)
    r = tf.paragraphs[0].add_run(); r.text = '✓  '
    r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = C['primary']
    r2 = tf.paragraphs[0].add_run(); r2.text = check
    r2.font.size = Pt(12); r2.font.color.rgb = C['text']
    y += 0.55
# CTA
ctf = tb(s, 0.6, 6.2, 12.1, 0.6)
p(ctf, 'Comece: npm install -g g360-cli  |  g360 init meu-projeto --template lit-web', 13, True, C['primary'], align=PP_ALIGN.CENTER)

prs.save('manual-g360-cli.pptx')
print(f'[OK] Manual gerado: manual-g360-cli.pptx ({TOTAL} slides)')
